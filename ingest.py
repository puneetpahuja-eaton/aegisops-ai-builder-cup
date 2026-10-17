"""Document ingestion: extract -> chunk -> embed -> store as per-device RAG corpus.

Supports PDF, Word (.docx), Excel (.xlsx/.csv) and plain text. Extracted text is
chunked, embedded (Vertex AI or offline fallback) and stored under the device's
`chunks` subcollection for retrieval-augmented generation.
"""
from __future__ import annotations

import datetime
import io
import os
import re
import uuid

import datastore
import devices
from embeddings import embed_texts

SOURCE_KINDS = (
    "datasheet", "dataset", "testcase", "fielddata", "eventlog", "powerlog",
    "document", "github", "confluence", "manual", "other",
)
_CHUNK_CHARS = 3500
_OVERLAP_CHARS = 400
_MAX_CHARS = 2_000_000


def _now() -> str:
    return datetime.datetime.now(datetime.timezone.utc).isoformat()


# --------------------------------------------------------------------------- #
# Extraction
# --------------------------------------------------------------------------- #
def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _extract_pdf(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith(".xlsx"):
        return _extract_xlsx(data)
    if name.endswith(".csv"):
        return data.decode("utf-8-sig", errors="replace")
    if name.endswith((".txt", ".md", ".json", ".yaml", ".yml", ".py", ".log")):
        return data.decode("utf-8", errors="replace")
    # Best effort for unknown types.
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [para.text for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Chunking
# --------------------------------------------------------------------------- #
def chunk_text(text: str, size: int = _CHUNK_CHARS, overlap: int = _OVERLAP_CHARS) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", (text or "").strip())
    if not text:
        return []
    if len(text) <= size:
        return [text]
    chunks, start = [], 0
    while start < len(text):
        end = start + size
        chunk = text[start:end]
        # Prefer to break on a paragraph/sentence boundary near the end.
        boundary = max(chunk.rfind("\n\n"), chunk.rfind(". "))
        if boundary > size // 2 and end < len(text):
            chunk = chunk[: boundary + 1]
            end = start + len(chunk)
        chunks.append(chunk.strip())
        start = max(end - overlap, end) if end >= len(text) else end - overlap
    return [chunk for chunk in chunks if chunk]


# --------------------------------------------------------------------------- #
# Optional raw-file archival to GCS
# --------------------------------------------------------------------------- #
def _archive_to_gcs(device_id: str, filename: str, data: bytes) -> str | None:
    bucket_name = os.getenv("GCS_BUCKET")
    if not bucket_name:
        return None
    from google.cloud import storage as gcs

    credentials = None
    if os.getenv("GOOGLE_OAUTH_ACCESS_TOKEN"):
        from google.oauth2.credentials import Credentials
        credentials = Credentials(os.environ["GOOGLE_OAUTH_ACCESS_TOKEN"])
    client = gcs.Client(project=os.environ.get("GOOGLE_CLOUD_PROJECT"), credentials=credentials)
    prefix = os.getenv("GCS_PREFIX", "")
    blob_name = f"{prefix}uploads/{device_id}/{uuid.uuid4().hex}_{filename}"
    client.bucket(bucket_name).blob(blob_name).upload_from_string(data)
    return f"gs://{bucket_name}/{blob_name}"


# --------------------------------------------------------------------------- #
# Ingestion
# --------------------------------------------------------------------------- #
def ingest_text(device_id: str, *, name: str, kind: str, text: str, gcs_uri: str | None = None,
                link_id: str | None = None) -> dict:
    if not devices.get_device(device_id):
        raise ValueError("No such device.")
    if kind not in SOURCE_KINDS:
        kind = "other"
    if len(text) > _MAX_CHARS:
        text = text[:_MAX_CHARS]
    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("No extractable text found in this source.")
    embeddings = embed_texts(chunks)
    source_id = uuid.uuid4().hex
    chunks_collection = devices.chunks_path(device_id)
    for index, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
        datastore.set_doc(chunks_collection, f"{source_id}-{index}", {
            "source_id": source_id, "source_name": name, "kind": kind, "link_id": link_id,
            "index": index, "text": chunk, "embedding": embedding,
        })
    record = {
        "id": source_id, "name": name, "kind": kind, "gcs_uri": gcs_uri, "link_id": link_id,
        "chunk_count": len(chunks), "char_count": len(text), "ingested_at": _now(), "status": "ready",
    }
    datastore.set_doc(devices.sources_path(device_id), source_id, record)
    devices.update_device(device_id, {"status": "ready"})
    return record


def ingest_file(device_id: str, filename: str, data: bytes, kind: str = "document") -> dict:
    text = extract_text(filename, data)
    gcs_uri = None
    try:
        gcs_uri = _archive_to_gcs(device_id, filename, data)
    except Exception:
        gcs_uri = None  # Archival is best-effort; ingestion still proceeds.
    return ingest_text(device_id, name=filename, kind=kind, text=text, gcs_uri=gcs_uri)


def list_sources(device_id: str) -> list[dict]:
    return sorted(datastore.list_docs(devices.sources_path(device_id)),
                  key=lambda s: s.get("ingested_at", ""))


def delete_source(device_id: str, source_id: str) -> None:
    chunks_collection = devices.chunks_path(device_id)
    for chunk in datastore.list_docs(chunks_collection):
        if chunk.get("source_id") == source_id:
            datastore.delete_doc(chunks_collection, chunk["id"])
    datastore.delete_doc(devices.sources_path(device_id), source_id)


def delete_by_link(device_id: str, link_id: str) -> int:
    """Remove all sources and chunks that came from one knowledge link (for refresh)."""
    removed = 0
    for source in datastore.list_docs(devices.sources_path(device_id)):
        if source.get("link_id") == link_id:
            delete_source(device_id, source["id"])
            removed += 1
    return removed


def corpus_size(device_id: str) -> int:
    return datastore.count_docs(devices.chunks_path(device_id))
