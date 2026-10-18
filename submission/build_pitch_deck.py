"""Build the Kairos AI pitch deck (6 slides, PPTX) for the AI Builder Cup.

Generates `Kairos-AI-Pitch-Deck.pptx`; a companion step converts it to PDF via
PowerPoint. Six slides cover: cover, problem+solution, five-agent engine,
Google Cloud & Gen AI stack (with high-level architecture + where/how mapping
and the rules-compliance check), impact (downtime/cost + social), and vision.

All Google Cloud usage is grounded in the code:
  - Gemini 2.5 Flash on Vertex AI        (gemini_client.py, vertexai=True)
  - Vertex AI embeddings text-embedding-004 (embeddings.py)
  - Firestore                             (datastore.py)
  - Cloud Storage                         (storage.py / device_data.py)
  - BigQuery                              (bq.py)
  - Secret Manager                        (secretmgr.py)
  - Cloud Run deployment                  (cloud/deploy.ps1, service 'kairos-ai')
Quantified savings are labeled illustrative/synthetic.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent

# ---- palette -------------------------------------------------------------- #
BG = RGBColor(0x0B, 0x1B, 0x2E)
PANEL = RGBColor(0x14, 0x2C, 0x47)
PANEL2 = RGBColor(0x1C, 0x3A, 0x5C)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
BLUE = RGBColor(0x4C, 0x8D, 0xFF)
AMBER = RGBColor(0xF5, 0xA5, 0x24)
GREEN = RGBColor(0x34, 0xD3, 0x99)
PINK = RGBColor(0xF4, 0x72, 0xB6)
WHITE = RGBColor(0xF4, 0xF7, 0xFB)
MUTE = RGBColor(0x9F, 0xB3, 0xC8)
DARKTEXT = RGBColor(0x0B, 0x1B, 0x2E)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SH = 7.5


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def _no_shadow(shape):
    shape.shadow.inherit = False


def box(s, l, t, w, h, fill=PANEL, line=None, line_w=1.0, radius=0.07, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape_type, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    _no_shadow(sp)
    try:
        if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = radius
    except Exception:
        pass
    sp.text_frame.word_wrap = True
    return sp


def text(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, r in enumerate(runs):
        txt, size, bold, color = r[0], r[1], r[2], r[3]
        space_after = r[4] if len(r) > 4 else 4
        space_before = r[5] if len(r) > 5 else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(space_before)
        run = p.add_run()
        run.text = txt
        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.name = FONT
        f.color.rgb = color
    return tb


def fill_text(shape, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    for i, r in enumerate(runs):
        txt, size, bold, color = r[0], r[1], r[2], r[3]
        space_after = r[4] if len(r) > 4 else 3
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = txt
        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.name = FONT
        f.color.rgb = color


def header(s, kicker, title, accent=TEAL):
    box(s, 0.6, 0.42, 0.14, 0.6, fill=accent, radius=0.3)
    text(s, 0.85, 0.36, 11.8, 0.35, [(kicker.upper(), 12, True, accent)])
    text(s, 0.85, 0.6, 11.8, 0.65, [(title, 26, True, WHITE)])


def footer(s, n):
    text(s, 0.6, 7.05, 9.5, 0.33, [("Kairos AI  ·  Google Cloud AI Builder Cup 2026  ·  Manufacturing", 9, False, MUTE)])
    text(s, 11.5, 7.05, 1.25, 0.33, [(f"{n} / 6", 9, True, MUTE)], align=PP_ALIGN.RIGHT)


def arrow(s, l, t, w, h, fill=TEAL, shape_type=MSO_SHAPE.RIGHT_ARROW):
    sp = s.shapes.add_shape(shape_type, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    _no_shadow(sp)
    return sp


def chip(s, l, t, w, h, label, fill=PANEL, color=WHITE, size=11, bold=True, line=None):
    sp = box(s, l, t, w, h, fill=fill, line=line, radius=0.3)
    fill_text(sp, [(label, size, bold, color)], align=PP_ALIGN.CENTER)
    return sp


# =========================================================================== #
# Slide 1 — Cover
# =========================================================================== #
s = slide()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), Inches(SH))
band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background(); _no_shadow(band)
circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.6), Inches(-2.2), Inches(6.0), Inches(6.0))
circ.fill.solid(); circ.fill.fore_color.rgb = PANEL; circ.line.fill.background(); _no_shadow(circ)
circ2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.0), Inches(3.4), Inches(4.2), Inches(4.2))
circ2.fill.solid(); circ2.fill.fore_color.rgb = PANEL2; circ2.line.fill.background(); _no_shadow(circ2)
text(s, 0.95, 1.45, 11.5, 0.5, [("GOOGLE CLOUD AI BUILDER CUP 2026  ·  JAPAC", 13, True, TEAL)])
text(s, 0.9, 1.95, 11.6, 1.5, [("Kairos AI", 70, True, WHITE)])
text(s, 0.95, 3.4, 10.8, 0.9, [("An expert reliability engineer inside every power device.", 22, True, TEAL)])
text(s, 0.95, 4.25, 10.5, 1.0,
     [("Autonomous, remote fault diagnosis and self-healing for critical power infrastructure "
       "— catching the moment to act, before downtime.", 15, False, MUTE)])
chip(s, 0.95, 5.4, 3.3, 0.5, "Theme · Manufacturing", fill=PANEL, size=12)
chip(s, 4.4, 5.4, 3.5, 0.5, "Gemini 2.5 Flash · Vertex AI", fill=PANEL, size=12)
chip(s, 8.1, 5.4, 3.2, 0.5, "Deployed on Cloud Run", fill=PANEL, size=12)
text(s, 0.95, 6.35, 11.5, 0.5,
     [("Live: kairos-ai-498775210896.asia-south1.run.app     ·     github.com/puneetpahuja-eaton/kairos-ai",
       11, False, MUTE)])

# =========================================================================== #
# Slide 2 — Problem & Solution
# =========================================================================== #
s = slide()
header(s, "The problem — and our answer", "Today, finding a fault needs a human on site", accent=AMBER)
# Problem panel
lp = box(s, 0.6, 1.5, 5.95, 3.7, fill=PANEL)
box(s, 0.6, 1.5, 5.95, 0.12, fill=AMBER, radius=0.3)
text(s, 0.85, 1.72, 5.5, 0.4, [("THE PROBLEM", 12.5, True, AMBER)])
text(s, 0.85, 2.12, 5.5, 1.2,
     [("A fault is only caught when someone notices an alarm, reports it, and waits for an engineer to "
       "travel to site. Downtime compounds — and expert engineers are scarce and can't be everywhere.",
       13, False, WHITE)])
for i, (h, b) in enumerate([
        ("No root-cause intelligence", "Tools only alert; they don't explain why."),
        ("No prediction or correlation", "Cascading alarms hide the true cause."),
        ("No closed-loop remediation", "Only ~10% (alerting) is automated; 90% is manual.")]):
    t = 3.45 + i * 0.58
    text(s, 0.95, t, 5.4, 0.55, [(f"•  {h} — ", 12, True, WHITE, 0), (b, 12, False, MUTE)])
# Solution panel
rp = box(s, 6.75, 1.5, 5.95, 3.7, fill=PANEL)
box(s, 6.75, 1.5, 5.95, 0.12, fill=TEAL, radius=0.3)
text(s, 7.0, 1.72, 5.5, 0.4, [("KAIROS AI", 12.5, True, TEAL)])
text(s, 7.0, 2.12, 5.5, 0.9,
     [("Puts an expert reliability engineer inside every device — learning it from its own documents, "
       "watching it remotely 24/7, and closing the loop from detection to verified fix.", 13, False, WHITE)])
for i, (h, b, acc) in enumerate([
        ("Learn", "Ingests datasheets & logs (RAG on Vertex AI).", TEAL),
        ("Monitor", "Watches live telemetry & event logs remotely.", BLUE),
        ("Diagnose", "Plain-English root cause, cited from the docs.", GREEN),
        ("Act & Verify", "Self-heals or escalates — human-approved.", AMBER)]):
    t = 3.25 + i * 0.47
    text(s, 7.1, t, 5.3, 0.45, [(f"{h}  ", 12.5, True, acc, 0), ("— " + b, 12, False, MUTE)])
# device strip
strip = box(s, 0.6, 5.45, 12.1, 1.25, fill=PANEL2)
text(s, 0.85, 5.6, 11.8, 0.35, [("APPLIES TO ANY DEVICE THAT EXPOSES TELEMETRY OR LOGS", 11.5, True, TEAL)])
text(s, 0.85, 5.98, 11.9, 0.35,
     [("Panelboards · Meters · Switchgear · UPS & batteries · PDUs · Transformers · Motors · Breakers",
       13, True, WHITE)])
text(s, 0.85, 6.33, 11.9, 0.35,
     [("Protecting: hospitals · water & utilities · telecom · manufacturing · data centers · substations",
       12, False, MUTE)])
footer(s, 2)

# =========================================================================== #
# Slide 3 — Five-agent engine
# =========================================================================== #
s = slide()
header(s, "How it works", "A closed loop of five agents, guided by Gemini")
gp = box(s, 0.6, 1.4, 12.1, 0.95, fill=PANEL2, line=TEAL, line_w=1.25)
fill_text(gp, [("Gemini Knowledge & Agentic Plane  ·  Gemini 2.5 Flash on Vertex AI", 14, True, TEAL),
               ("RAG retrieval over device docs   ·   Ask / Plan / Agent assistant   ·   dataset & live-monitor-plan generators",
                11, False, WHITE)], align=PP_ALIGN.CENTER)
arrow(s, 5.72, 2.37, 0.35, 0.4, fill=TEAL, shape_type=MSO_SHAPE.DOWN_ARROW)
agents = [
    ("1 · Monitor", "Detects anomalies in live & historic telemetry (robust-Z + Isolation Forest).", TEAL),
    ("2 · Diagnose", "Root cause via topology correlation + Gemini hypotheses with cited evidence.", BLUE),
    ("3 · Decide", "Confidence-gated autonomy — weighs reversibility & asset criticality.", GREEN),
    ("4 · Act", "Self-heals reversible faults or hands off to a human.", AMBER),
    ("5 · Verify", "Confirms the fix; scores downtime & cost avoided.", PINK),
]
aw, ah = 2.2, 2.55
xs = [0.65, 3.1, 5.55, 8.0, 10.45]
for i, (h, b, acc) in enumerate(agents):
    l = xs[i]
    box(s, l, 2.95, aw, ah, fill=PANEL, line=acc, line_w=1.25)
    box(s, l, 2.95, aw, 0.12, fill=acc, radius=0.3)
    text(s, l + 0.18, 3.15, aw - 0.36, 0.5, [(h, 13.5, True, acc)])
    text(s, l + 0.18, 3.68, aw - 0.36, 1.7, [(b, 11, False, WHITE)])
    if i < 4:
        arrow(s, l + aw - 0.02, 4.0, 0.28, 0.42, fill=MUTE)
hil = box(s, 5.55, 5.65, 4.65, 0.5, fill=PANEL2, line=AMBER, line_w=1.0)
fill_text(hil, [("Human-in-the-loop approval at Decide → Act", 11, True, AMBER)], align=PP_ALIGN.CENTER)
loop = box(s, 0.65, 6.25, 12.0, 0.5, fill=None, line=TEAL, line_w=1.25, radius=0.2)
fill_text(loop, [("↺  Continuous closed loop — Verify feeds Monitor, so the engine keeps learning  ·  one engine for live + historic",
                  11, True, TEAL)], align=PP_ALIGN.CENTER)
footer(s, 3)

# =========================================================================== #
# Slide 4 — Google Cloud & Gen AI stack (architecture + where/how + compliance)
# =========================================================================== #
s = slide()
header(s, "Google Cloud & Gen AI stack", "Built end-to-end on Google Cloud", accent=BLUE)
# architecture flow (compact band)
dchip = box(s, 0.6, 1.4, 1.75, 0.95, fill=PANEL2, line=TEAL, line_w=1.0)
fill_text(dchip, [("Field devices", 11.5, True, WHITE), ("REST + event logs", 9.5, False, MUTE)], align=PP_ALIGN.CENTER)
arrow(s, 2.42, 1.75, 0.35, 0.3, fill=TEAL)
crun = box(s, 2.85, 1.4, 7.6, 0.95, fill=PANEL, line=BLUE, line_w=1.0)
fill_text(crun, [("Kairos AI on Cloud Run  —  Gemini · Embeddings · Firestore · Cloud Storage · BigQuery · Secret Manager",
                  11.5, True, WHITE)], align=PP_ALIGN.CENTER)
arrow(s, 10.52, 1.75, 0.35, 0.3, fill=TEAL)
echip = box(s, 10.95, 1.4, 1.75, 0.95, fill=PANEL2, line=GREEN, line_w=1.0)
fill_text(echip, [("Engineer", 11.5, True, WHITE), ("dashboard · chat", 9.5, False, MUTE)], align=PP_ALIGN.CENTER)
# mapping table
rows = [
    ("Google Cloud service", "Where in Kairos AI", "How it is used"),
    ("Gemini 2.5 Flash · Vertex AI", "Ask-the-device, Diagnose agent, generators", "Grounded RCA, plain-English answers, dataset & monitor-plan generation"),
    ("Vertex AI Embeddings (text-embedding-004)", "Data Sources · RAG retrieval", "Embeds each device's docs into a per-device knowledge base"),
    ("Firestore", "Platform state", "Users, devices, sources, chunks & knowledge links — durable across instances"),
    ("Cloud Storage", "Telemetry & datasets", "Stores generated/uploaded datasets, SOPs and archives"),
    ("BigQuery", "Audit trail", "Partitioned audit table with 30-day retention"),
    ("Secret Manager", "Integrations", "GitHub / Confluence keys — never stored in device records"),
    ("Cloud Run", "Deployment", "Serverless host; dedicated service account via ADC; scales to zero"),
]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(2.6), Inches(12.1), Inches(3.55)).table
tbl.columns[0].width = Inches(3.5)
tbl.columns[1].width = Inches(3.5)
tbl.columns[2].width = Inches(5.1)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.1); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
            col, bold, size = WHITE, True, 11.5
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = PANEL if ri % 2 else BG
            col, bold, size = WHITE, (ci == 0), 10.5
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = val
        run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = (TEAL if (ci == 0 and ri > 0) else col); run.font.name = FONT
# compliance strip
comp = box(s, 0.6, 6.3, 12.1, 0.52, fill=PANEL2, line=GREEN, line_w=1.0)
fill_text(comp, [("RULES CHECK   ✅ Google AI model: Gemini 2.5 Flash    ✅ Deployed on Cloud Run    "
                  "✅ Agentic 5-agent engine    ✅ English · functional · scalable", 11.5, True, GREEN)],
          align=PP_ALIGN.CENTER)
footer(s, 4)

# =========================================================================== #
# Slide 5 — Impact (downtime/cost + social)
# =========================================================================== #
s = slide()
header(s, "Impact", "Hours of downtime → minutes, and lives kept on", accent=GREEN)
rows = [
    ("Fault (asset)", "Path", "Down", "$ avoided"),
    ("Voltage sag (Main Incomer)", "Auto-resolved", "40m", "$34,000"),
    ("Overload (Feeder PDU-2)", "Auto-resolved", "25m", "$11,250"),
    ("Battery degradation (UPS-3)", "Escalated (early)", "48m", "$40,800"),
    ("Thermal runaway (UPS-1)", "Auto-resolved", "55m", "$46,750"),
    ("Harmonic distortion (PDU-1)", "Auto-resolved", "35m", "$15,750"),
    ("Breaker trip (BRK-4)", "Auto-resolved", "20m", "$5,000"),
    ("Total · 6 incidents", "5 auto / 1 esc.", "223m", "$153,550"),
]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.55), Inches(6.5), Inches(4.0)).table
tbl.columns[0].width = Inches(3.1)
tbl.columns[1].width = Inches(1.7)
tbl.columns[2].width = Inches(0.8)
tbl.columns[3].width = Inches(0.9)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.05)
        cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = GREEN
            col, bold, size = DARKTEXT, True, 11
        elif ri == len(rows) - 1:
            cell.fill.solid(); cell.fill.fore_color.rgb = PANEL2
            col, bold, size = WHITE, True, 11
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = PANEL if ri % 2 else BG
            col, bold, size = WHITE, False, 10.5
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = val
        run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = col; run.font.name = FONT
text(s, 0.6, 5.65, 6.5, 1.0,
     [("Illustrative synthetic scenarios — not realized savings. Rates: critical $850/min, high $450/min, "
       "medium $250/min. MTTR drops from tens of hours to 1–2 cycles.", 9.5, False, MUTE)])
# social impact cards (right)
text(s, 7.35, 1.5, 5.4, 0.4, [("WHY IT MATTERS", 12.5, True, PINK)])
impacts = [
    ("Essential services uptime", "Hospitals, water & telecom stay running.", GREEN),
    ("Expertise for everyone", "Senior-grade diagnosis without a specialist on site.", BLUE),
    ("Safety", "Early detection of fire/overload hazards prevents accidents.", AMBER),
    ("Sustainability", "Cuts wasted energy, extends equipment life, fewer truck-rolls.", TEAL),
]
for i, (h, b, acc) in enumerate(impacts):
    t = 2.0 + i * 1.18
    box(s, 7.35, t, 5.4, 1.02, fill=PANEL)
    box(s, 7.35, t, 0.1, 1.02, fill=acc, radius=0.3)
    text(s, 7.62, t + 0.14, 5.0, 0.4, [(h, 14, True, WHITE)])
    text(s, 7.62, t + 0.52, 5.0, 0.45, [(b, 11.5, False, MUTE)])
footer(s, 5)

# =========================================================================== #
# Slide 6 — Vision / closing
# =========================================================================== #
s = slide()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), Inches(SH))
band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background(); _no_shadow(band)
circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.8), Inches(3.0), Inches(5.5), Inches(5.5))
circ.fill.solid(); circ.fill.fore_color.rgb = PANEL; circ.line.fill.background(); _no_shadow(circ)
text(s, 0.95, 1.4, 11.2, 0.5, [("WHY KAIROS AI WINS", 13, True, TEAL)])
for i, (h, b, acc) in enumerate([
        ("Deep, real Google Cloud AI", "Gemini on Vertex AI, embeddings, Firestore, Cloud Storage, BigQuery, Secret Manager, Cloud Run.", TEAL),
        ("One engine, live + historic", "The same five-agent brain diagnoses a dataset and a live stream.", BLUE),
        ("The agent decides what to watch", "It reads the datasheet to pick parameters and decode event logs.", GREEN),
        ("Honest by design", "Offline fallbacks labeled; demo data synthetic; humans approve actions.", AMBER)]):
    t = 1.95 + i * 0.6
    text(s, 0.95, t, 11.3, 0.55, [(f"✓  {h}  ", 13.5, True, acc, 0), ("— " + b, 12, False, MUTE)])
text(s, 0.9, 4.55, 11.4, 1.3, [("Start with one device.", 38, True, WHITE), ("Scale to the world.", 38, True, TEAL)])
text(s, 0.95, 6.15, 11.3, 0.5,
     [("Live · kairos-ai-498775210896.asia-south1.run.app      ·      github.com/puneetpahuja-eaton/kairos-ai",
       11.5, False, MUTE)])
text(s, 0.95, 6.6, 11.3, 0.4, [("121 automated tests · deployed & verified on Cloud Run · Firestore-backed · Gemini on Vertex AI",
                                10.5, True, TEAL)])

out = HERE / "Kairos-AI-Pitch-Deck.pptx"
prs.save(str(out))
print(f"Saved {out}  ({len(prs.slides._sldIdLst)} slides)")
